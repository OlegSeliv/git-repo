"""
Модуль для работы с Excel таблицами
Создание отчетов и таблиц учета
"""

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from datetime import datetime
from database import db
import os


class ExcelManager:
    def __init__(self, output_dir: str = "excel_reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
    
    def create_payment_tracking_table(self, contract_id: int) -> str:
        """
        Создать таблицу учета оплат против выполненных работ
        """
        # Получить данные
        contract = db.execute_query("SELECT * FROM contracts WHERE id = ?", (contract_id,))[0]
        orders = db.execute_query(
            "SELECT * FROM work_orders WHERE contract_id = ? ORDER BY created_at",
            (contract_id,)
        )
        payments = db.execute_query(
            "SELECT * FROM payments WHERE contract_id = ? ORDER BY payment_date",
            (contract_id,)
        )
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Учет оплат"
        
        # Стили
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Заголовок документа
        ws['A1'] = f'Учет оплат и выполненных работ по контракту {contract["contract_number"]}'
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:H1')
        
        ws['A2'] = f'Заказчик: {contract["customer"]}'
        ws['A3'] = f'Период: {contract["start_date"]} - {contract["end_date"]}'
        ws['A4'] = f'Сумма контракта: {contract["total_amount"]} руб.'
        
        # Таблица заказов
        row = 6
        ws[f'A{row}'] = 'ВЫПОЛНЕННЫЕ РАБОТЫ'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        
        row += 1
        headers = ['№ заказа', 'Устройство', 'Местоположение', 'Плановая стоимость', 'Фактическая стоимость', 'Статус', 'Срок']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        row += 1
        for order in orders:
            ws.cell(row=row, column=1, value=order['order_number']).border = border
            ws.cell(row=row, column=2, value=order['device_name']).border = border
            ws.cell(row=row, column=3, value=order['device_location'] or 'На заводе').border = border
            ws.cell(row=row, column=4, value=order['estimated_cost'] or 0).border = border
            ws.cell(row=row, column=5, value=order['actual_cost'] or 0).border = border
            ws.cell(row=row, column=6, value=order['status']).border = border
            ws.cell(row=row, column=7, value=order['repair_deadline']).border = border
            row += 1
        
        # Итого по работам
        total_estimated = sum(o['estimated_cost'] or 0 for o in orders)
        total_actual = sum(o['actual_cost'] or 0 for o in orders)
        ws.cell(row=row, column=3, value='ИТОГО:').font = Font(bold=True)
        ws.cell(row=row, column=4, value=total_estimated).font = Font(bold=True)
        ws.cell(row=row, column=5, value=total_actual).font = Font(bold=True)
        
        # Таблица платежей
        row += 3
        ws[f'A{row}'] = 'ПЛАТЕЖИ'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        
        row += 1
        payment_headers = ['Дата', 'Тип платежа', 'Этап', 'Сумма', 'Статус', 'Номер счета']
        for col, header in enumerate(payment_headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        row += 1
        for payment in payments:
            ws.cell(row=row, column=1, value=payment['payment_date']).border = border
            ws.cell(row=row, column=2, value=payment['payment_type']).border = border
            ws.cell(row=row, column=3, value=payment['payment_stage']).border = border
            ws.cell(row=row, column=4, value=payment['amount']).border = border
            ws.cell(row=row, column=5, value=payment['status']).border = border
            ws.cell(row=row, column=6, value=payment['invoice_number']).border = border
            row += 1
        
        # Итого по платежам
        total_paid = sum(p['amount'] for p in payments if p['status'] == 'completed')
        total_expected = sum(p['amount'] for p in payments if p['status'] == 'pending')
        
        ws.cell(row=row, column=3, value='ИТОГО оплачено:').font = Font(bold=True)
        ws.cell(row=row, column=4, value=total_paid).font = Font(bold=True)
        row += 1
        ws.cell(row=row, column=3, value='Ожидается:').font = Font(bold=True)
        ws.cell(row=row, column=4, value=total_expected).font = Font(bold=True)
        
        # Баланс
        row += 2
        ws[f'A{row}'] = 'ФИНАНСОВЫЙ БАЛАНС'
        ws[f'A{row}'].font = Font(bold=True, size=12)
        row += 1
        
        balance_data = [
            ['Сумма контракта', contract['total_amount']],
            ['Фактически потрачено', total_actual],
            ['Получено оплат', total_paid],
            ['Ожидается оплат', total_expected],
            ['Доступный остаток', total_paid - total_actual],
            ['Ожидаемый остаток', total_paid + total_expected - total_actual]
        ]
        
        for label, value in balance_data:
            ws.cell(row=row, column=1, value=label).font = Font(bold=True)
            cell = ws.cell(row=row, column=2, value=value)
            cell.font = Font(bold=True if 'остаток' in label.lower() else False)
            if 'остаток' in label.lower():
                cell.fill = PatternFill(start_color="D9EAD3" if value >= 0 else "F4CCCC", 
                                       end_color="D9EAD3" if value >= 0 else "F4CCCC", 
                                       fill_type="solid")
            row += 1
        
        # Автоподбор ширины столбцов
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Сохранить
        filename = f'payment_tracking_{contract["contract_number"]}_{datetime.now().strftime("%Y%m%d")}.xlsx'
        filepath = os.path.join(self.output_dir, filename)
        wb.save(filepath)
        
        return filepath
    
    def create_components_deficit_report(self, contract_id: int = None) -> str:
        """
        Создать отчет по дефициту комплектующих
        """
        query = """
            SELECT 
                c.component_code,
                c.component_name,
                c.warehouse_stock,
                c.unit_of_measure,
                c.unit_price,
                SUM(oc.required_quantity) as total_required,
                SUM(oc.allocated_from_warehouse + oc.received_from_supplier) as total_allocated,
                SUM(oc.required_quantity - oc.allocated_from_warehouse - oc.received_from_supplier) as deficit,
                GROUP_CONCAT(DISTINCT wo.order_number) as affected_orders,
                GROUP_CONCAT(DISTINCT co.contract_number) as affected_contracts
            FROM order_components oc
            JOIN components c ON oc.component_id = c.id
            JOIN work_orders wo ON oc.work_order_id = wo.id
            JOIN contracts co ON wo.contract_id = co.id
        """
        params = []
        if contract_id:
            query += " WHERE wo.contract_id = ?"
            params.append(contract_id)
        
        query += """
            GROUP BY c.id
            HAVING deficit > 0
            ORDER BY deficit DESC
        """
        
        deficit_data = db.execute_query(query, tuple(params))
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Дефицит комплектующих"
        
        # Заголовок
        ws['A1'] = 'ОТЧЕТ ПО ДЕФИЦИТУ КОМПЛЕКТУЮЩИХ'
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:J1')
        ws['A2'] = f'Дата формирования: {datetime.now().strftime("%d.%m.%Y %H:%M")}'
        
        # Заголовки таблицы
        row = 4
        headers = [
            'Код', 'Наименование', 'Требуется', 'На складе', 'Выделено', 
            'ДЕФИЦИТ', 'Ед. изм.', 'Цена за ед.', 'Сумма дефицита', 'Затронутые заказы'
        ]
        
        header_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Данные
        row += 1
        total_deficit_cost = 0
        
        for item in deficit_data:
            deficit_cost = item['deficit'] * (item['unit_price'] or 0)
            total_deficit_cost += deficit_cost
            
            ws.cell(row=row, column=1, value=item['component_code'])
            ws.cell(row=row, column=2, value=item['component_name'])
            ws.cell(row=row, column=3, value=item['total_required'])
            ws.cell(row=row, column=4, value=item['warehouse_stock'])
            ws.cell(row=row, column=5, value=item['total_allocated'])
            
            deficit_cell = ws.cell(row=row, column=6, value=item['deficit'])
            deficit_cell.font = Font(bold=True, color="FF0000")
            deficit_cell.fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
            
            ws.cell(row=row, column=7, value=item['unit_of_measure'])
            ws.cell(row=row, column=8, value=item['unit_price'])
            ws.cell(row=row, column=9, value=deficit_cost)
            ws.cell(row=row, column=10, value=item['affected_orders'])
            
            row += 1
        
        # Итого
        ws.cell(row=row, column=5, value='ИТОГО стоимость дефицита:').font = Font(bold=True)
        ws.cell(row=row, column=9, value=total_deficit_cost).font = Font(bold=True, size=12)
        
        # Автоподбор ширины
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Сохранить
        filename = f'components_deficit_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx'
        filepath = os.path.join(self.output_dir, filename)
        wb.save(filepath)
        
        return filepath
    
    def create_trips_schedule(self, contract_id: int = None) -> str:
        """
        Создать расписание командировок
        """
        query = """
            SELECT 
                bt.*,
                e.full_name,
                e.position,
                c.contract_number,
                wo.order_number,
                wo.device_location
            FROM business_trips bt
            JOIN employees e ON bt.employee_id = e.id
            JOIN contracts c ON bt.contract_id = c.id
            LEFT JOIN work_orders wo ON bt.work_order_id = wo.id
            WHERE bt.status IN ('planned', 'active')
        """
        params = []
        if contract_id:
            query += " AND bt.contract_id = ?"
            params.append(contract_id)
        
        query += " ORDER BY bt.start_date"
        
        trips = db.execute_query(query, tuple(params))
        
        wb = Workbook()
        ws = wb.active
        ws.title = "График командировок"
        
        # Заголовок
        ws['A1'] = 'ГРАФИК КОМАНДИРОВОК'
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:I1')
        
        # Таблица
        row = 3
        headers = [
            'ФИО', 'Должность', 'Направление', 'Начало', 'Окончание', 
            'Дней', 'Контракт', 'Заказ', 'Статус'
        ]
        
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        
        # Данные
        row += 1
        for trip in trips:
            start = datetime.strptime(trip['start_date'], "%Y-%m-%d")
            end = datetime.strptime(trip['planned_end_date'], "%Y-%m-%d")
            days = (end - start).days
            
            ws.cell(row=row, column=1, value=trip['full_name'])
            ws.cell(row=row, column=2, value=trip['position'])
            ws.cell(row=row, column=3, value=trip['destination'])
            ws.cell(row=row, column=4, value=trip['start_date'])
            ws.cell(row=row, column=5, value=trip['planned_end_date'])
            ws.cell(row=row, column=6, value=days)
            ws.cell(row=row, column=7, value=trip['contract_number'])
            ws.cell(row=row, column=8, value=trip['order_number'])
            ws.cell(row=row, column=9, value=trip['status'])
            
            row += 1
        
        # Автоподбор
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Сохранить
        filename = f'trips_schedule_{datetime.now().strftime("%Y%m%d")}.xlsx'
        filepath = os.path.join(self.output_dir, filename)
        wb.save(filepath)
        
        return filepath


# Создание экземпляра
excel_manager = ExcelManager()
