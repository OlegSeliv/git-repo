"""
Модуль генерации документов
"""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime, date
from pathlib import Path
import json
from database import DatabaseManager

class DocumentGenerator:
    def __init__(self):
        self.db = DatabaseManager()
        self.templates_dir = Path("templates")
        self.output_dir = Path("generated_documents")
        self.templates_dir.mkdir(exist_ok=True)
        self.output_dir.mkdir(exist_ok=True)
    
    def generate_completion_act(self, order_id: int) -> str:
        """Генерация акта выполненных работ"""
        # Получение данных
        order_query = '''
        SELECT o.*, c.contract_number, c.customer_name
        FROM orders o
        JOIN contracts c ON o.contract_id = c.id
        WHERE o.id = ?
        '''
        order_data = self.db.execute_query(order_query, (order_id,))[0]
        
        # Создание документа
        doc = Document()
        
        # Заголовок
        title = doc.add_heading('АКТ СДАЧИ-ПРИЕМКИ ВЫПОЛНЕННЫХ РАБОТ', level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Номер и дата
        doc.add_paragraph(f'№ {order_data["order_number"]} от {datetime.now().strftime("%d.%m.%Y")}')
        doc.add_paragraph()
        
        # Информация о контракте
        doc.add_paragraph(f'По контракту № {order_data["contract_number"]}')
        doc.add_paragraph(f'Заказчик: {order_data["customer_name"]}')
        doc.add_paragraph()
        
        # Описание работ
        doc.add_heading('Выполненные работы:', level=2)
        doc.add_paragraph(order_data["description"])
        
        # Место выполнения
        doc.add_paragraph(f'Место выполнения работ: {order_data["location"]}')
        
        if order_data["device_serial"]:
            doc.add_paragraph(f'Серийный номер устройства: {order_data["device_serial"]}')
        
        # Сроки
        doc.add_paragraph()
        doc.add_paragraph(f'Начало работ: {order_data.get("actual_start", "не указано")}')
        doc.add_paragraph(f'Окончание работ: {order_data.get("actual_end", "не указано")}')
        
        # Стоимость
        doc.add_paragraph()
        doc.add_paragraph(f'Стоимость работ: {order_data.get("actual_cost", 0):.2f} руб.')
        
        # Заключение
        doc.add_paragraph()
        doc.add_paragraph('Работы выполнены в полном объеме, в соответствии с условиями контракта.')
        doc.add_paragraph('Претензий по качеству и срокам выполнения работ не имеется.')
        
        # Подписи
        doc.add_paragraph()
        doc.add_paragraph()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Исполнитель:'
        table.cell(0, 1).text = 'Заказчик:'
        table.cell(1, 0).text = '_______________ /_______________/'
        table.cell(1, 1).text = '_______________ /_______________/'
        
        # Сохранение
        filename = f'act_completion_{order_id}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx'
        file_path = self.output_dir / filename
        doc.save(file_path)
        
        # Сохранение в БД
        query = '''
        INSERT INTO documents (
            document_type, order_id, title, file_path, status
        ) VALUES ('completion_act', ?, ?, ?, 'generated')
        '''
        self.db.execute_update(query, (order_id, f'Акт выполненных работ {order_data["order_number"]}', str(file_path)))
        
        return str(file_path)
    
    def generate_deficiency_letter(self, contract_id: int) -> str:
        """Генерация письма о нехватке средств"""
        # Получение данных
        contract = self.db.get_contract_financial_status(contract_id)
        
        doc = Document()
        
        # Заголовок
        doc.add_paragraph(f'Исх. № _____ от {datetime.now().strftime("%d.%m.%Y")}')
        doc.add_paragraph()
        doc.add_paragraph('Кому: ' + self.db.execute_query(
            "SELECT customer_name FROM contracts WHERE id = ?", (contract_id,)
        )[0]['customer_name'])
        doc.add_paragraph()
        
        title = doc.add_heading('ПИСЬМО', level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading('о необходимости дополнительного финансирования', level=2)
        
        # Основной текст
        doc.add_paragraph()
        doc.add_paragraph(f'Уважаемые коллеги!')
        doc.add_paragraph()
        doc.add_paragraph(f'В рамках выполнения работ по контракту № {contract["contract_number"]} '
                         f'информируем Вас о следующем:')
        
        doc.add_paragraph()
        doc.add_paragraph(f'1. Общая стоимость контракта: {contract["total_amount"]:.2f} руб.')
        doc.add_paragraph(f'2. Получено средств: {contract["total_received"]:.2f} руб.')
        doc.add_paragraph(f'3. Фактические расходы: {contract["total_expenses"]:.2f} руб.')
        doc.add_paragraph(f'4. Текущий баланс: {contract["balance"]:.2f} руб.')
        
        if contract["balance"] < 0:
            doc.add_paragraph()
            doc.add_paragraph(f'Для завершения работ по контракту требуется дополнительное '
                            f'финансирование в размере {abs(contract["balance"]):.2f} руб.')
        
        doc.add_paragraph()
        doc.add_paragraph('Просим рассмотреть возможность выделения дополнительных средств '
                         'для успешного завершения работ в установленные сроки.')
        
        # Подпись
        doc.add_paragraph()
        doc.add_paragraph()
        doc.add_paragraph('С уважением,')
        doc.add_paragraph('Руководитель проекта')
        doc.add_paragraph('_______________ /_______________/')
        
        # Сохранение
        filename = f'deficiency_letter_{contract_id}_{datetime.now().strftime("%Y%m%d_%H%M%S")}.docx'
        file_path = self.output_dir / filename
        doc.save(file_path)
        
        return str(file_path)
    
    def generate_materials_report(self, order_id: int = None) -> str:
        """Генерация отчета по материалам"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Отчет по материалам"
        
        # Заголовки
        headers = ['Артикул', 'Наименование', 'Ед.изм.', 'На складе', 
                  'Зарезервировано', 'Доступно', 'Минимум', 'Статус', 'Поставщик']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Данные
        if order_id:
            materials = self.db.check_material_availability(order_id)
        else:
            query = '''
            SELECT 
                article, name, unit, quantity_warehouse,
                quantity_reserved,
                (quantity_warehouse - quantity_reserved) as available,
                quantity_minimum,
                CASE 
                    WHEN (quantity_warehouse - quantity_reserved) < quantity_minimum 
                    THEN 'Дефицит'
                    ELSE 'В наличии'
                END as status,
                supplier
            FROM materials
            ORDER BY name
            '''
            materials = self.db.execute_query(query)
        
        for row_num, material in enumerate(materials, 2):
            ws.cell(row=row_num, column=1, value=material.get('article'))
            ws.cell(row=row_num, column=2, value=material.get('name'))
            ws.cell(row=row_num, column=3, value=material.get('unit', 'шт'))
            ws.cell(row=row_num, column=4, value=material.get('quantity_warehouse', 0))
            ws.cell(row=row_num, column=5, value=material.get('quantity_reserved', 0))
            ws.cell(row=row_num, column=6, value=material.get('available', 0))
            ws.cell(row=row_num, column=7, value=material.get('quantity_minimum', 0))
            
            status_cell = ws.cell(row=row_num, column=8, value=material.get('status'))
            if material.get('status') == 'Дефицит' or material.get('status') == 'deficit':
                status_cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            
            ws.cell(row=row_num, column=9, value=material.get('supplier', ''))
        
        # Автоширина колонок
        for column_cells in ws.columns:
            length = max(len(str(cell.value or '')) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = length + 2
        
        # Сохранение
        filename = f'materials_report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        file_path = self.output_dir / filename
        wb.save(file_path)
        
        return str(file_path)
    
    def generate_business_trips_schedule(self) -> str:
        """Генерация графика командировок"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "График командировок"
        
        # Заголовки
        headers = ['Сотрудник', 'Место', 'Начало', 'Окончание', 
                  'Продлено до', 'Заказ', 'Контракт', 'Статус', 'Расходы план', 'Расходы факт']
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        
        # Данные
        trips = self.db.get_active_business_trips()
        
        for row_num, trip in enumerate(trips, 2):
            ws.cell(row=row_num, column=1, value=trip['employee_name'])
            ws.cell(row=row_num, column=2, value=trip['location'])
            ws.cell(row=row_num, column=3, value=trip['start_date'])
            ws.cell(row=row_num, column=4, value=trip['end_date'])
            ws.cell(row=row_num, column=5, value=trip.get('extended_to', ''))
            ws.cell(row=row_num, column=6, value=trip.get('order_number', ''))
            ws.cell(row=row_num, column=7, value=trip.get('contract_number', ''))
            
            status_cell = ws.cell(row=row_num, column=8, value=trip['status'])
            if trip['status'] == 'active':
                status_cell.fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")
            
            ws.cell(row=row_num, column=9, value=trip.get('estimated_expenses', 0))
            ws.cell(row=row_num, column=10, value=trip.get('actual_expenses', 0))
        
        # Автоширина колонок
        for column_cells in ws.columns:
            length = max(len(str(cell.value or '')) for cell in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = length + 2
        
        # Сохранение
        filename = f'business_trips_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        file_path = self.output_dir / filename
        wb.save(file_path)
        
        return str(file_path)